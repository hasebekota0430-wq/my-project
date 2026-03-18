"""人事制度運用フォーマット生成スクリプト"""
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ─────────────────────────────────────────────
# 共通スタイル定義
# ─────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
SUB_FILL      = PatternFill("solid", fgColor="2E75B6")
ALT_FILL      = PatternFill("solid", fgColor="D6E4F0")
SECTION_FILL  = PatternFill("solid", fgColor="BDD7EE")
YELLOW_FILL   = PatternFill("solid", fgColor="FFF2CC")
GREEN_FILL    = PatternFill("solid", fgColor="E2EFDA")
RED_FILL      = PatternFill("solid", fgColor="FCE4D6")
WHITE_FILL    = PatternFill("solid", fgColor="FFFFFF")

H_FONT  = Font(name="游ゴシック", bold=True, color="FFFFFF", size=11)
S_FONT  = Font(name="游ゴシック", bold=True, color="FFFFFF", size=10)
B_FONT  = Font(name="游ゴシック", bold=True, size=10)
N_FONT  = Font(name="游ゴシック", size=10)
T_FONT  = Font(name="游ゴシック", bold=True, size=14, color="1F4E79")

THIN = Side(style="thin", color="4472C4")
MED  = Side(style="medium", color="1F4E79")
thin_border  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
thick_border = Border(left=MED,  right=MED,  top=MED,  bottom=MED)

def hdr(ws, row, col, value, fill=HEADER_FILL, font=H_FONT, wrap=True):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = fill; c.font = font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    c.border = thin_border
    return c

def cell(ws, row, col, value="", fill=WHITE_FILL, font=N_FONT,
         align="left", wrap=True, num_fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = fill; c.font = font
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = thin_border
    if num_fmt:
        c.number_format = num_fmt
    return c

def title_row(ws, row, text, col_span=(1, 10)):
    c = ws.cell(row=row, column=col_span[0], value=text)
    c.font = T_FONT
    c.fill = PatternFill("solid", fgColor="DEEAF1")
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = thick_border
    ws.merge_cells(start_row=row, start_column=col_span[0],
                   end_row=row,   end_column=col_span[1])
    ws.row_dimensions[row].height = 28

def merge_hdr(ws, r1, r2, c1, c2, value, fill=HEADER_FILL, font=H_FONT):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    c = ws.cell(row=r1, column=c1, value=value)
    c.fill = fill; c.font = font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border


# ══════════════════════════════════════════════
# Sheet 1: 目次
# ══════════════════════════════════════════════
ws0 = wb.active
ws0.title = "目次"
ws0.column_dimensions["A"].width = 6
ws0.column_dimensions["B"].width = 36
ws0.column_dimensions["C"].width = 50
ws0.column_dimensions["D"].width = 16

title_row(ws0, 1, "人事制度運用フォーマット　目次", (1, 4))
hdr(ws0, 2, 1, "No.", SUB_FILL, S_FONT)
hdr(ws0, 2, 2, "シート名", SUB_FILL, S_FONT)
hdr(ws0, 2, 3, "用途・概要", SUB_FILL, S_FONT)
hdr(ws0, 2, 4, "更新頻度", SUB_FILL, S_FONT)

sheets_info = [
    (1, "①等級・職位マスタ",   "等級定義・職位一覧・昇格要件",       "年1回以上"),
    (2, "②社員台帳",           "全社員の基本情報・等級・雇用情報",   "随時"),
    (3, "③目標管理（MBO）",    "期初目標設定／期末評価記録",         "半期ごと"),
    (4, "④人事評価シート",     "一次評価・二次評価・最終評価",       "半期ごと"),
    (5, "⑤昇格・降格管理",     "昇降格申請・審査・承認フロー",       "随時"),
    (6, "⑥給与・報酬管理",     "基本給・手当・賞与の管理",           "月次/年次"),
    (7, "⑦採用管理",           "採用計画・選考状況・内定管理",       "随時"),
    (8, "⑧研修・育成管理",     "受講記録・スキルマップ",             "随時"),
    (9, "⑨入退社管理",         "入社手続き・退社手続きチェック",     "随時"),
]

for i, (no, name, desc, freq) in enumerate(sheets_info, start=3):
    f = ALT_FILL if i % 2 == 0 else WHITE_FILL
    cell(ws0, i, 1, no,   fill=f, align="center")
    cell(ws0, i, 2, name, fill=f, font=B_FONT)
    cell(ws0, i, 3, desc, fill=f)
    cell(ws0, i, 4, freq, fill=f, align="center")

ws0.row_dimensions[1].height = 32
ws0.row_dimensions[2].height = 22
for r in range(3, 13):
    ws0.row_dimensions[r].height = 20

cell(ws0, 13, 1, "※ 本フォーマットは叩き台です。貴社制度に合わせて適宜修正してください。",
     fill=YELLOW_FILL, font=Font(name="游ゴシック", size=9, italic=True, color="7F7F7F"))
ws0.merge_cells("A13:D13")


# ══════════════════════════════════════════════
# Sheet 2: 等級・職位マスタ
# ══════════════════════════════════════════════
ws1 = wb.create_sheet("①等級・職位マスタ")
cols = [6, 18, 18, 14, 42, 16, 16]
for i, w in enumerate(cols, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

title_row(ws1, 1, "等級・職位マスタ", (1, 7))

headers = ["等級コード", "等級名称", "職位名称", "役職区分", "職務・役割の定義（概要）", "給与レンジ下限", "給与レンジ上限"]
for col, h in enumerate(headers, 1):
    hdr(ws1, 2, col, h)

grades = [
    ("G1", "スタッフ1級",   "一般社員",     "一般職", "指示のもと定型業務を遂行する",             "200,000", "250,000"),
    ("G2", "スタッフ2級",   "シニアスタッフ","一般職", "自律的に定型業務を遂行・後輩指導を担う",    "250,000", "310,000"),
    ("G3", "リーダー1級",   "チームリーダー","管理職", "小チームをマネジメントし業務推進を図る",    "310,000", "380,000"),
    ("G4", "リーダー2級",   "課長代理",     "管理職", "課内の目標達成・人材育成を推進する",        "380,000", "450,000"),
    ("G5", "マネジャー1級", "課長",         "管理職", "課の方針策定・組織運営・P&L管理",           "450,000", "540,000"),
    ("G6", "マネジャー2級", "部長",         "管理職", "部門戦略立案・経営方針への貢献",            "540,000", "680,000"),
    ("G7", "エグゼクティブ","本部長/役員",  "経営職", "全社戦略の策定・意思決定・対外折衝",        "680,000", "—"),
]

for i, row in enumerate(grades, start=3):
    f = ALT_FILL if i % 2 == 0 else WHITE_FILL
    for col, val in enumerate(row, 1):
        cell(ws1, i, col, val, fill=f, align="center" if col in (1,4,6,7) else "left")

ws1.row_dimensions[2].height = 22
for r in range(3, 11):
    ws1.row_dimensions[r].height = 20

cell(ws1, 11, 1, "※ 給与レンジは月額（円）。等級数・名称は自社制度に合わせて変更してください。",
     fill=YELLOW_FILL, font=Font(name="游ゴシック", size=9, italic=True, color="7F7F7F"))
ws1.merge_cells("A11:G11")


# ══════════════════════════════════════════════
# Sheet 3: 社員台帳
# ══════════════════════════════════════════════
ws2 = wb.create_sheet("②社員台帳")
col_widths = [10,14,14,8,10,16,16,10,10,10,10,14,18,16,12]
for i, w in enumerate(col_widths, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

title_row(ws2, 1, "社員台帳", (1, 15))

h2 = ["社員番号","姓","名","性別","生年月日","所属部署","職位名称","等級",
      "雇用区分","入社年月日","在籍年数","直属上長ID","緊急連絡先","メールアドレス","備考"]
for col, h in enumerate(h2, 1):
    hdr(ws2, 2, col, h)

sample = [
    ["EMP001","山田","太郎","男","1985/04/15","営業部","課長","G5",
     "正社員","2010/04/01","=DATEDIF(J3,TODAY(),\"Y\")&\"年\"","EMP020","090-XXXX-XXXX","yamada@example.com",""],
    ["EMP002","佐藤","花子","女","1992/08/20","人事部","シニアスタッフ","G2",
     "正社員","2016/04/01","=DATEDIF(J4,TODAY(),\"Y\")&\"年\"","EMP015","090-XXXX-XXXX","sato@example.com","育休復帰"],
    ["EMP003","鈴木","一郎","男","1990/11/03","開発部","チームリーダー","G3",
     "正社員","2014/04/01","=DATEDIF(J5,TODAY(),\"Y\")&\"年\"","EMP018","090-XXXX-XXXX","suzuki@example.com",""],
]

for i, row in enumerate(sample, start=3):
    f = ALT_FILL if i % 2 == 0 else WHITE_FILL
    for col, val in enumerate(row, 1):
        c = cell(ws2, i, col, val if col != 11 else None, fill=f,
                 align="center" if col in (1,4,8,9,10,11) else "left")
        if col == 11:
            try:
                c.value = val
            except:
                c.value = ""

for r in range(3, 8):
    ws2.row_dimensions[r].height = 20

cell(ws2, 8, 1, "※ 個人情報の取り扱いには十分注意し、アクセス権限を適切に設定してください。",
     fill=RED_FILL, font=Font(name="游ゴシック", size=9, bold=True, color="C00000"))
ws2.merge_cells("A8:O8")


# ══════════════════════════════════════════════
# Sheet 4: 目標管理（MBO）
# ══════════════════════════════════════════════
ws3 = wb.create_sheet("③目標管理（MBO）")
col_widths3 = [10,14,10,10,36,10,14,10,12,12,14,20]
for i, w in enumerate(col_widths3, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

title_row(ws3, 1, "目標管理シート（MBO）", (1, 12))

# サブヘッダー：評価期間
cell(ws3, 2, 1, "評価期間：", fill=SECTION_FILL, font=B_FONT, align="right")
cell(ws3, 2, 2, "2025年度 上期", fill=YELLOW_FILL, font=B_FONT)
ws3.merge_cells("B2:D2")
cell(ws3, 2, 5, "社員番号：", fill=SECTION_FILL, font=B_FONT, align="right")
cell(ws3, 2, 6, "", fill=YELLOW_FILL)
cell(ws3, 2, 7, "氏名：", fill=SECTION_FILL, font=B_FONT, align="right")
cell(ws3, 2, 8, "", fill=YELLOW_FILL)
ws3.merge_cells("H2:I2")
cell(ws3, 2, 10, "所属：", fill=SECTION_FILL, font=B_FONT, align="right")
cell(ws3, 2, 11, "", fill=YELLOW_FILL)
ws3.merge_cells("K2:L2")

headers3_1 = ["No.","目標カテゴリ","重み\n(%)", "難易度",
              "目標内容（具体的に記載）","KPI\n目標値", "KPI\n実績値",
              "達成率\n(%)", "自己評価\n(S/A/B/C)", "一次評価\n(S/A/B/C)",
              "評価コメント（上長）","次期への申し送り事項"]
for col, h in enumerate(headers3_1, 1):
    hdr(ws3, 3, col, h)

mbo_rows = [
    ["1","業績目標","40%","A","新規顧客獲得 XX件／月","5件","","","","","",""],
    ["2","業績目標","20%","B","既存顧客売上維持率 XX%以上","95%","","","","","",""],
    ["3","組織目標","20%","B","チーム残業時間削減（月XX時間以内）","20h","","","","","",""],
    ["4","能力開発","10%","B","資格取得：XX試験 合格","合格","","","","","",""],
    ["5","その他","10%","C","社内勉強会 企画・運営（XX回）","2回","","","","","",""],
    ["","合計","100%","","","","","=IFERROR(AVERAGE(H4:H8),\"\")","","","",""],
]

for i, row in enumerate(mbo_rows, start=4):
    f = ALT_FILL if (i-4) % 2 == 0 else WHITE_FILL
    if i == 9:
        f = GREEN_FILL
    for col, val in enumerate(row, 1):
        c = cell(ws3, i, col, val if not val.startswith("=") else None,
                 fill=f, align="center" if col in (1,3,4,6,7,8,9,10) else "left")
        if val.startswith("="):
            c.value = val

cell(ws3, 10, 1, "【総合評価欄】", fill=SECTION_FILL, font=B_FONT)
ws3.merge_cells("A10:B10")
cell(ws3, 10, 3, "総合自己評価：", fill=SECTION_FILL, font=B_FONT)
ws3.merge_cells("C10:D10")
cell(ws3, 10, 5, "", fill=YELLOW_FILL)
cell(ws3, 10, 6, "総合一次評価：", fill=SECTION_FILL, font=B_FONT)
ws3.merge_cells("F10:G10")
cell(ws3, 10, 8, "", fill=YELLOW_FILL)
cell(ws3, 10, 9, "総合最終評価：", fill=SECTION_FILL, font=B_FONT)
ws3.merge_cells("I10:J10")
cell(ws3, 10, 11, "", fill=YELLOW_FILL)
ws3.merge_cells("K10:L10")

cell(ws3, 11, 1, "【本人コメント】", fill=SECTION_FILL, font=B_FONT)
ws3.merge_cells("A11:B11")
cell(ws3, 11, 3, "", fill=WHITE_FILL)
ws3.merge_cells("C11:L11")
ws3.row_dimensions[11].height = 50

cell(ws3, 12, 1, "評価基準：S=120%以上　A=100%以上　B=80%以上　C=80%未満",
     fill=YELLOW_FILL, font=Font(name="游ゴシック", size=9, italic=True))
ws3.merge_cells("A12:L12")

ws3.row_dimensions[3].height = 36
for r in range(4, 12):
    ws3.row_dimensions[r].height = 22


# ══════════════════════════════════════════════
# Sheet 5: 人事評価シート
# ══════════════════════════════════════════════
ws4 = wb.create_sheet("④人事評価シート")
col_widths4 = [10,14,10,36,10,10,10,12,12,12,20]
for i, w in enumerate(col_widths4, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w

title_row(ws4, 1, "人事評価シート", (1, 11))

# 基本情報
for col, (label, width) in enumerate([("社員番号",8),("氏名",14),("等級",6),
                                        ("所属",14),("評価期間",10),("評価者",12)], start=1):
    hdr(ws4, 2, col, label, SUB_FILL, S_FONT)
for col in range(1, 7):
    cell(ws4, 3, col, "", fill=YELLOW_FILL)

# 評価ヘッダー
eval_cats = [
    ("業績評価", "目標達成度・成果物の質・量"),
    ("プロセス評価", "取り組み姿勢・協働・改善提案"),
    ("コンピテンシー評価", "リーダーシップ・専門知識・コミュニケーション"),
    ("行動規範評価", "コンプライアンス・誠実さ・顧客志向"),
]

row = 4
hdr(ws4, row, 1, "評価区分", HEADER_FILL, H_FONT)
hdr(ws4, row, 2, "評価項目", HEADER_FILL, H_FONT)
ws4.merge_cells(f"B{row}:C{row}")
hdr(ws4, row, 4, "評価基準・着眼点", HEADER_FILL, H_FONT)
hdr(ws4, row, 5, "重み\n(%)", HEADER_FILL, H_FONT)
hdr(ws4, row, 6, "自己評価\n(1-5)", HEADER_FILL, H_FONT)
hdr(ws4, row, 7, "一次評価\n(1-5)", HEADER_FILL, H_FONT)
hdr(ws4, row, 8, "二次評価\n(1-5)", HEADER_FILL, H_FONT)
hdr(ws4, row, 9, "最終評価\n(1-5)", HEADER_FILL, H_FONT)
hdr(ws4, row, 10, "加重点", HEADER_FILL, H_FONT)
hdr(ws4, row, 11, "コメント", HEADER_FILL, H_FONT)

eval_items = [
    ("業績評価",            "目標達成度",       "設定目標に対する達成率・成果の質",    "40%"),
    ("業績評価",            "生産性・効率性",   "同等成果を効率よく生み出せているか",  "10%"),
    ("プロセス評価",        "積極性・自律性",   "指示待ちでなく自ら動いているか",      "10%"),
    ("プロセス評価",        "チームワーク",     "協力・情報共有・後輩支援",            "10%"),
    ("コンピテンシー評価",  "専門知識・スキル", "職務遂行に必要な知識・技術の習得度",  "15%"),
    ("コンピテンシー評価",  "課題解決力",       "問題の発見・分析・解決策の実行",      "5%"),
    ("行動規範評価",        "コンプライアンス", "法令・社内規定の遵守",                "5%"),
    ("行動規範評価",        "顧客志向",         "顧客・関係者の期待に応える行動",      "5%"),
]

row = 5
for i, (cat, item, criterion, weight) in enumerate(eval_items, start=row):
    f = ALT_FILL if (i - row) % 2 == 0 else WHITE_FILL
    cell(ws4, i, 1, cat,       fill=f, align="center")
    cell(ws4, i, 2, item,      fill=f)
    ws4.merge_cells(f"B{i}:C{i}")
    cell(ws4, i, 4, criterion, fill=f)
    cell(ws4, i, 5, weight,    fill=f, align="center")
    cell(ws4, i, 6, "",        fill=YELLOW_FILL, align="center")
    cell(ws4, i, 7, "",        fill=YELLOW_FILL, align="center")
    cell(ws4, i, 8, "",        fill=YELLOW_FILL, align="center")
    cell(ws4, i, 9, "",        fill=YELLOW_FILL, align="center")
    cell(ws4, i, 10, f"=IF(I{i}<>\"\",I{i}*VALUE(SUBSTITUTE(E{i},\"%\",\"\"))/100,\"\")",
         fill=GREEN_FILL, align="center")
    cell(ws4, i, 11, "",       fill=WHITE_FILL)

# 合計行
tot = row + len(eval_items)
hdr(ws4, tot, 1, "合計", SUB_FILL, S_FONT)
ws4.merge_cells(f"A{tot}:E{tot}")
cell(ws4, tot, 6, f"=IFERROR(AVERAGE(F{row}:F{tot-1}),\"\")", fill=GREEN_FILL, align="center")
cell(ws4, tot, 7, f"=IFERROR(AVERAGE(G{row}:G{tot-1}),\"\")", fill=GREEN_FILL, align="center")
cell(ws4, tot, 8, f"=IFERROR(AVERAGE(H{row}:H{tot-1}),\"\")", fill=GREEN_FILL, align="center")
cell(ws4, tot, 9, f"=IFERROR(AVERAGE(I{row}:I{tot-1}),\"\")", fill=GREEN_FILL, align="center")
cell(ws4, tot, 10, f"=IFERROR(SUM(J{row}:J{tot-1}),\"\")",   fill=GREEN_FILL, align="center", font=B_FONT)

cell(ws4, tot+1, 1, "評価基準：5=卓越　4=優秀　3=標準　2=改善要　1=不十分",
     fill=YELLOW_FILL, font=Font(name="游ゴシック", size=9, italic=True))
ws4.merge_cells(f"A{tot+1}:K{tot+1}")

ws4.row_dimensions[4].height = 36
for r in range(5, tot+2):
    ws4.row_dimensions[r].height = 22


# ══════════════════════════════════════════════
# Sheet 6: 昇格・降格管理
# ══════════════════════════════════════════════
ws5 = wb.create_sheet("⑤昇格・降格管理")
col_widths5 = [10,14,10,12,12,12,12,14,10,10,20,20]
for i, w in enumerate(col_widths5, 1):
    ws5.column_dimensions[get_column_letter(i)].width = w

title_row(ws5, 1, "昇格・降格管理台帳", (1, 12))
headers5 = ["申請日","氏名","社員番号","現等級","変更後等級",
            "変更区分","変更予定日","申請者（上長）","一次承認","最終承認","推薦理由","備考"]
for col, h in enumerate(headers5, 1):
    hdr(ws5, 2, col, h)

sample5 = [
    ["2025/04/01","山田 太郎","EMP001","G4","G5","昇格","2025/04/01","部長 田中","承認","承認","目標達成率120%、リーダーシップ発揮",""],
    ["2025/04/01","佐藤 花子","EMP002","G1","G2","昇格","2025/04/01","課長 鈴木","承認","承認","業務習熟・後輩指導実績あり",""],
]
for i, row in enumerate(sample5, start=3):
    f = ALT_FILL if i % 2 == 0 else WHITE_FILL
    for col, val in enumerate(row, 1):
        cell(ws5, i, col, val, fill=f,
             align="center" if col in (1,3,4,5,6,7,9,10) else "left")

for r in range(3, 8):
    ws5.row_dimensions[r].height = 20

cell(ws5, 8, 1, "変更区分選択肢：昇格 / 降格 / 等級変更（横異動）/ その他",
     fill=YELLOW_FILL, font=Font(name="游ゴシック", size=9, italic=True))
ws5.merge_cells("A8:L8")


# ══════════════════════════════════════════════
# Sheet 7: 給与・報酬管理
# ══════════════════════════════════════════════
ws6 = wb.create_sheet("⑥給与・報酬管理")
col_widths6 = [10,14,10,12,12,12,12,12,12,12,12,14]
for i, w in enumerate(col_widths6, 1):
    ws6.column_dimensions[get_column_letter(i)].width = w

title_row(ws6, 1, "給与・報酬管理台帳", (1, 12))
headers6 = ["社員番号","氏名","等級","基本給","職務手当","通勤手当",
            "その他手当","総支給額","社保控除","所得税","差引支給額","最終更新日"]
for col, h in enumerate(headers6, 1):
    hdr(ws6, 2, col, h)

sample6 = [
    ["EMP001","山田 太郎","G5",450000,50000,15000,0,"=SUM(D3:G3)",80000,"=ROUND((H3-I3)*0.1,0)","=H3-I3-J3","2025/04/01"],
    ["EMP002","佐藤 花子","G2",260000,20000,12000,0,"=SUM(D4:G4)",50000,"=ROUND((H4-I4)*0.1,0)","=H4-I4-J4","2025/04/01"],
    ["EMP003","鈴木 一郎","G3",320000,30000,10000,0,"=SUM(D5:G5)",60000,"=ROUND((H5-I5)*0.1,0)","=H5-I5-J5","2025/04/01"],
]
for i, row in enumerate(sample6, start=3):
    f = ALT_FILL if i % 2 == 0 else WHITE_FILL
    for col, val in enumerate(row, 1):
        c = cell(ws6, i, col, val if not str(val).startswith("=") else None,
                 fill=f, align="center" if col in (1,3,12) else "right" if col >= 4 else "left")
        if str(val).startswith("="):
            c.value = val
        if col in range(4, 12) and col != 12:
            c.number_format = '#,##0'

cell(ws6, 8, 1, "※ 社保・税額は簡易計算。実際は給与計算システムと連携してください。",
     fill=RED_FILL, font=Font(name="游ゴシック", size=9, bold=True, color="C00000"))
ws6.merge_cells("A8:L8")
for r in range(3, 8):
    ws6.row_dimensions[r].height = 20


# ══════════════════════════════════════════════
# Sheet 8: 採用管理
# ══════════════════════════════════════════════
ws7 = wb.create_sheet("⑦採用管理")
col_widths7 = [10,16,12,10,10,10,10,10,10,12,14,14,14,16]
for i, w in enumerate(col_widths7, 1):
    ws7.column_dimensions[get_column_letter(i)].width = w

title_row(ws7, 1, "採用管理台帳", (1, 14))
headers7 = ["応募番号","氏名","応募職種","応募経路","応募日",
            "書類選考","一次面接","二次面接","最終面接","内定日",
            "内定承諾","入社予定日","担当採用者","備考・特記事項"]
for col, h in enumerate(headers7, 1):
    hdr(ws7, 2, col, h)

sample7 = [
    ["R2025-001","候補者 A","エンジニア","自社HP","2025/03/01","通過","通過","通過","通過","2025/03/20","承諾","2025/04/01","採用担当 田中",""],
    ["R2025-002","候補者 B","営業","転職エージェント","2025/03/05","通過","通過","通過","—","—","—","—","採用担当 山本","最終面接調整中"],
    ["R2025-003","候補者 C","デザイナー","リファラル","2025/03/10","通過","—","—","—","—","—","—","採用担当 田中","一次面接設定中"],
]
status_colors = {"通過": GREEN_FILL, "承諾": GREEN_FILL, "—": PatternFill("solid", fgColor="F2F2F2")}
for i, row in enumerate(sample7, start=3):
    for col, val in enumerate(row, 1):
        f = status_colors.get(val, WHITE_FILL) if col in range(6, 12) else (ALT_FILL if i%2==0 else WHITE_FILL)
        cell(ws7, i, col, val, fill=f, align="center" if col != 14 else "left")

for r in range(3, 8):
    ws7.row_dimensions[r].height = 20
cell(ws7, 8, 1, "応募経路選択肢：自社HP / 転職エージェント / リファラル / 求人サイト / 新卒媒体 / その他",
     fill=YELLOW_FILL, font=Font(name="游ゴシック", size=9, italic=True))
ws7.merge_cells("A8:N8")


# ══════════════════════════════════════════════
# Sheet 9: 研修・育成管理
# ══════════════════════════════════════════════
ws8 = wb.create_sheet("⑧研修・育成管理")
col_widths8 = [10,14,10,20,12,12,8,10,10,10,10,16]
for i, w in enumerate(col_widths8, 1):
    ws8.column_dimensions[get_column_letter(i)].width = w

title_row(ws8, 1, "研修・育成管理台帳", (1, 12))
headers8 = ["社員番号","氏名","等級","研修名称","研修分類",
            "実施日","時間数","受講形態","合否/修了","取得スキル","費用(円)","備考"]
for col, h in enumerate(headers8, 1):
    hdr(ws8, 2, col, h)

sample8 = [
    ["EMP001","山田 太郎","G5","管理職研修 上級編","マネジメント","2025/02/10","8h","集合","修了","リーダーシップ","50,000","外部研修"],
    ["EMP002","佐藤 花子","G2","Excel上級講座","ITスキル","2025/01/20","4h","オンライン","修了","表計算","5,000",""],
    ["EMP003","鈴木 一郎","G3","プロジェクトマネジメント","専門","2025/03/05","16h","集合","受講中","PM基礎","30,000","資格試験準備"],
]
for i, row in enumerate(sample8, start=3):
    f = ALT_FILL if i % 2 == 0 else WHITE_FILL
    for col, val in enumerate(row, 1):
        cell(ws8, i, col, val, fill=f,
             align="center" if col in (1,3,5,6,7,8,9,11) else "left")

for r in range(3, 8):
    ws8.row_dimensions[r].height = 20
cell(ws8, 8, 1, "受講形態選択肢：集合 / オンライン / OJT / 自己学習 / eラーニング",
     fill=YELLOW_FILL, font=Font(name="游ゴシック", size=9, italic=True))
ws8.merge_cells("A8:L8")


# ══════════════════════════════════════════════
# Sheet 10: 入退社管理
# ══════════════════════════════════════════════
ws9 = wb.create_sheet("⑨入退社管理")
col_widths9 = [10,14,10,10,12,12,10,12,12,12,12,20]
for i, w in enumerate(col_widths9, 1):
    ws9.column_dimensions[get_column_letter(i)].width = w

title_row(ws9, 1, "入退社管理チェックリスト", (1, 12))
hdr(ws9, 2, 1, "社員番号", SUB_FILL, S_FONT)
hdr(ws9, 2, 2, "氏名", SUB_FILL, S_FONT)
hdr(ws9, 2, 3, "区分", SUB_FILL, S_FONT)
hdr(ws9, 2, 4, "入退社日", SUB_FILL, S_FONT)
ws9.merge_cells("E2:F2"); hdr(ws9, 2, 5, "労務手続き", SUB_FILL, S_FONT)
ws9.merge_cells("G2:H2"); hdr(ws9, 2, 7, "IT・セキュリティ", SUB_FILL, S_FONT)
ws9.merge_cells("I2:J2"); hdr(ws9, 2, 9, "備品・資産", SUB_FILL, S_FONT)
ws9.merge_cells("K2:L2"); hdr(ws9, 2, 11, "担当者・備考", SUB_FILL, S_FONT)

sub_headers = ["雇用保険\n加入/脱退","社保\n手続き","PC\n払出/回収","アカウント\n設定/削除","社員証\n発行/返却","備品\n確認"]
for col, h in enumerate(sub_headers, 5):
    hdr(ws9, 3, col, h, PatternFill("solid", fgColor="4472C4"),
        Font(name="游ゴシック", bold=True, color="FFFFFF", size=9))

sample9 = [
    ["EMP004","新入 四郎","入社","2025/04/01","済","済","済","済","済","済","人事 山田","入社案内送付済"],
    ["EMP010","退職 十郎","退社","2025/03/31","済","済","済","処理中","返却済","未確認","人事 佐藤","離職票発行予定"],
]
status_c = {"済": GREEN_FILL, "処理中": YELLOW_FILL, "未確認": RED_FILL, "返却済": GREEN_FILL}
for i, row in enumerate(sample9, start=4):
    for col, val in enumerate(row[:4], 1):
        f = ALT_FILL if i%2==0 else WHITE_FILL
        cell(ws9, i, col, val, fill=f, align="center" if col in (1,3,4) else "left")
    for col, val in enumerate(row[4:10], 5):
        f = status_c.get(val, WHITE_FILL)
        cell(ws9, i, col, val, fill=f, align="center")
    cell(ws9, i, 11, row[10], fill=ALT_FILL if i%2==0 else WHITE_FILL)
    cell(ws9, i, 12, row[11], fill=ALT_FILL if i%2==0 else WHITE_FILL)

ws9.row_dimensions[2].height = 22
ws9.row_dimensions[3].height = 36
for r in range(4, 9):
    ws9.row_dimensions[r].height = 22
cell(ws9, 9, 1, "チェック状態選択肢：済 / 処理中 / 未着手 / 不要",
     fill=YELLOW_FILL, font=Font(name="游ゴシック", size=9, italic=True))
ws9.merge_cells("A9:L9")


# ──────────────────────────────────────────────
# 全シート共通：先頭行を固定
# ──────────────────────────────────────────────
for ws in wb.worksheets:
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:2"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1

# 保存
output_path = "/home/user/my-project/人事制度運用フォーマット_叩き台.xlsx"
wb.save(output_path)
print(f"✓ 保存完了: {output_path}")
